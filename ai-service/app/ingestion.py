from io import BytesIO
import zipfile
from pypdf import PdfReader
from docx import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter


# 用途：阻止旧任务覆盖新版本或恢复已删除文档。
# 输入：current：当前文档记录；job_version：任务版本。
# 输出：是否仍可发布的布尔值。
def can_publish(current, job_version):
    # 旧任务完成时，文档可能已更新或删除；只有仍然有效的同一版本可以发布索引。
    return bool(current and not current["deleted"] and current["version"] == job_version)


# 用途：按扩展名提取文本并限制文件体积和内容。
# 输入：filename：文件名；content：原始字节。
# 输出：清理后的文本；不支持或无可读文本时抛 ValueError。
def parse_document(filename, content):
    suffix = filename.rsplit(".", 1)[-1].lower()
    if suffix in ("txt", "md"):
        text = content.decode("utf-8-sig")
        if "\x00" in text:
            raise ValueError("文本编码无效")
    elif suffix == "pdf":
        # pypdf 读取 PDF 的文本层，不提供 OCR；纯扫描件需要先做文字识别。
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted or len(reader.pages) > 500:
            raise ValueError("不支持加密或超长PDF")
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == "docx":
        with zipfile.ZipFile(BytesIO(content)) as z:
            if sum(x.file_size for x in z.infolist()) > 50 * 1024 * 1024:
                raise ValueError("文档解压后过大")
        doc = Document(BytesIO(content))
        text = "\n".join(
            [p.text for p in doc.paragraphs]
            + [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
        )
    else:
        raise ValueError("不支持的格式")
    if not text.strip() or len(text) > 2_000_000:
        raise ValueError("无可提取文本或文本过大")
    return text.strip()


# 用途：按段落和句子边界生成有重叠的检索片段。
# 输入：text：解析得到的全文。
# 输出：文本片段列表；切分单位为字符。
def split_document(text):
    # LangChain 递归按段落、句子等边界切分；900/140 是字符数，不是 token 数。
    # 相邻片段保留 140 字符重叠，减少一句话或一段规则被截断后丢失上下文的情况。
    return RecursiveCharacterTextSplitter(
        chunk_size=900, chunk_overlap=140, separators=["\n\n", "\n", "。", "；", " ", ""]
    ).split_text(text)
